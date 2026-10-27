"""Per-program dependency search evidence and an append-only Excel history.

No mainframe I/O occurs here. LibrarySearch supplies actual observed attempts;
shared dependencies replay their original evidence with an explicit cache note.
An offline inventory rebuild must never manufacture remote search outcomes.
"""
from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from mainframe.zowe_report_io import atomic_save_workbook, report_lock
from mainframe.zowe_workspace import search_history_path

HEADERS = (
    "IDMS Program Name", "Type", "Name of searching", "Python Search Library Path",
    "Result", "Remarks", "Run ID", "Observed At (UTC)",
)
SHEET_NAME = "Dependency Search History"
HEADER_ROW = 4
TYPE_LABELS = {"COPYBOOK": "Copybook", "DCLGEN": "DCLGENs", "SUBSCHEMA": "Sub Schema"}
RESULTS = {"Found", "Not Found", "Error", "Not Searched"}
MAX_ROWS = 1_048_576


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


@dataclass(frozen=True)
class SearchAttempt:
    library_path: str
    result: str
    remarks: str
    observed_at: str = field(default_factory=utc_now)


def http_attempt(dataset: str, member: str, status: int, *, probe: bool = False) -> SearchAttempt:
    """Classify observed status without mistaking access/transport errors for absence."""
    location = f"{dataset}({member})"
    if status == 200:
        message = "Member probe succeeded; content not downloaded." if probe else "Member content retrieved."
        return SearchAttempt(location, "Found", f"HTTP 200. {message}")
    if status == 404:
        return SearchAttempt(location, "Not Found", "HTTP 404. Requested dataset/member was not found.")
    descriptions = {
        400: "Invalid request; verify dataset/member and service configuration.",
        401: "Authentication failed; member existence is not established.",
        403: "Access denied; member existence is not established.",
        408: "Request timeout; member existence is not established.",
        429: "Service rate limit; member existence is not established.",
    }
    message = descriptions.get(status, "Service/request failure; member existence is not established.")
    return SearchAttempt(location, "Error", f"HTTP {status}. {message}")


def exception_attempt(dataset: str, member: str, exc: Exception) -> SearchAttempt:
    """Record safe diagnostic categories, never raw response bodies or credentials."""
    import requests
    if isinstance(exc, requests.Timeout):
        detail = "Request timed out."
    elif isinstance(exc, requests.exceptions.SSLError):
        detail = "TLS/certificate verification failed."
    elif isinstance(exc, requests.ConnectionError):
        detail = "Connection failed."
    elif isinstance(exc, ValueError):
        detail = "Invalid or missing connection/library configuration."
    else:
        detail = "Lookup failed; consult the secured operational log."
    return SearchAttempt(f"{dataset}({member})", "Error",
                         f"{type(exc).__name__}: {detail} Member existence is not established.")


@dataclass(frozen=True)
class HistoryRow:
    program: str
    artifact_type: str
    member: str
    library_path: str
    result: str
    remarks: str
    run_id: str
    observed_at: str

    def values(self) -> tuple[str, ...]:
        return (self.program, self.artifact_type, self.member, self.library_path,
                self.result, self.remarks, self.run_id, self.observed_at)


@dataclass
class SearchHistoryRecorder:
    """Run-scoped association: one trace per program/type/member, even on cache reuse."""
    run_id: str = field(default_factory=lambda: str(uuid4()))
    rows: list[HistoryRow] = field(default_factory=list)
    _recorded: set[tuple[str, str, str]] = field(default_factory=set)

    def record(self, program: str, artifact: str, member: str,
               attempts: Iterable[SearchAttempt], *, cached: bool = False) -> None:
        program, member, artifact = program.strip().upper(), member.strip().upper(), artifact.upper()
        if artifact not in TYPE_LABELS or not member or not program:
            raise ValueError("History needs a program, member and supported dependency type")
        key = (program, artifact, member)
        if key in self._recorded:
            return
        trace = tuple(attempts)
        if not trace:
            trace = (SearchAttempt("", "Not Searched", "No library attempt evidence was supplied."),)
        for attempt in trace:
            note = attempt.remarks
            observed = attempt.observed_at
            if cached:
                note += (f" Reused this run's earlier search evidence ({observed}); "
                         "no additional mainframe request for this program.")
                observed = utc_now()
            self.rows.append(HistoryRow(program, TYPE_LABELS[artifact], member,
                                        attempt.library_path, attempt.result, note,
                                        self.run_id, observed))
        self._recorded.add(key)

    def issue(self, program: str, artifact: str, member: str,
              library_path: str, remarks: str) -> None:
        """A download can succeed yet fail local saving/analysis: retain both facts."""
        self.rows.append(HistoryRow(program.upper(), TYPE_LABELS[artifact], member.upper(),
                                    library_path, "Error", remarks, self.run_id, utc_now()))


def _validate(values: tuple[str, ...]) -> None:
    if len(values) != len(HEADERS) or any(not isinstance(v, str) for v in values):
        raise ValueError("Search history rows must contain eight text fields")
    if any(len(v) > 32767 for v in values):
        raise ValueError("Excel cell limit exceeded; history will not be truncated")
    if not values[0].strip() or not values[2].strip() or not values[6].strip():
        raise ValueError("History requires program, member and Run ID")
    if values[1] not in TYPE_LABELS.values() or values[4] not in RESULTS:
        raise ValueError("Unsupported dependency type or result in search history")
    observed = datetime.fromisoformat(values[7])
    if observed.utcoffset() is None or observed.utcoffset().total_seconds() != 0:
        raise ValueError("History timestamp must include UTC timezone")
    if values[4] in {"Found", "Not Found"} and not values[3]:
        raise ValueError("Found/Not Found needs the actual observed library/member path")


def _new_workbook() -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "IDMS | Dependency library search history"
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17324D")
    sheet.row_dimensions[1].height = 36
    sheet.merge_cells("A2:H2")
    sheet["A2"] = ("Append-only observations per program and attempted library. Run ID + UTC time identify history. "
                   "Cached evidence is labelled; Found is not a conversion approval.")
    sheet["A2"].font = Font(name="Calibri", size=10, color="626C76")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 32
    sheet.merge_cells("A3:H3")
    sheet["A3"] = "No subschema reference = no schema search row. Untried libraries are never reported as Not Found."
    sheet["A3"].font = Font(name="Calibri", size=10, color="626C76")
    sheet.row_dimensions[3].height = 22
    for c, title in enumerate(HEADERS, 1):
        cell = sheet.cell(HEADER_ROW, c, title)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[HEADER_ROW].height = 32
    for c, width in zip("ABCDEFGH", (24, 18, 24, 49, 19, 82, 39, 36)):
        sheet.column_dimensions[c].width = width
    sheet.freeze_panes = "D5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    return book


def write_search_history(rows: Iterable[HistoryRow], path: Path | None = None) -> tuple[Path, int]:
    """Append observations atomically; exact replay is idempotent, new runs append.

    Refuse corrupt schemas, formulas and Excel limits rather than losing history.
    A caller may publish an empty workbook without claiming any mainframe search.
    """
    path = Path(path) if path is not None else search_history_path()
    incoming = list(rows)
    for row in incoming:
        _validate(row.values())
    with report_lock(path):
        book = load_workbook(path) if path.exists() else _new_workbook()
        try:
            if book.sheetnames != [SHEET_NAME]:
                raise ValueError("Unexpected search history sheet; original retained")
            sheet = book[SHEET_NAME]
            if sheet.max_column != len(HEADERS) or tuple(
                sheet.cell(HEADER_ROW, c).value for c in range(1, 9)
            ) != HEADERS:
                raise ValueError("Unexpected search history columns; original retained")
            seen: set[tuple[str, ...]] = set()
            for cells in sheet.iter_rows(min_row=HEADER_ROW + 1, max_col=8):
                if not any(c.value is not None for c in cells):
                    raise ValueError("Unexpected blank row in search history; original retained")
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Formula in search history; original retained")
                values = tuple(str(c.value or "") for c in cells)
                _validate(values)
                if values in seen:
                    raise ValueError("Duplicate historical observation; original retained")
                seen.add(values)
            for row in incoming:
                values = row.values()
                if values in seen:
                    continue
                if sheet.max_row >= MAX_ROWS:
                    raise ValueError("Excel row limit reached; archive history before continuing")
                sheet.append(values)
                number = sheet.max_row
                for c in sheet[number]:
                    c.data_type = "s"  # Source names/errors are text, never Excel formulas.
                    c.font = Font(name="Calibri", size=11,
                                  color="177245" if c.column < 5 else "626C76")
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                state = sheet.cell(number, 5)
                state.font = Font(name="Calibri", size=11, bold=True, color="17324D")
                state.fill = PatternFill("solid", fgColor={
                    "Found": "E2F0D9", "Not Found": "FFF2CC",
                    "Error": "FCE4D6", "Not Searched": "E7E6E6",
                }[row.result])
                estimated_lines = max(len(row.remarks) // 82 + 1,
                                      len(row.library_path) // 49 + 1, 2)
                sheet.row_dimensions[number].height = min(240, 15 * estimated_lines + 10)
                seen.add(values)
            if sheet.max_row > HEADER_ROW:
                table_ref = f"A4:H{sheet.max_row}"
                if "DependencySearches" in sheet.tables:
                    sheet.tables["DependencySearches"].ref = table_ref
                    if sheet.tables["DependencySearches"].autoFilter is not None:
                        sheet.tables["DependencySearches"].autoFilter.ref = table_ref
                else:
                    table = Table(displayName="DependencySearches", ref=table_ref)
                    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
                    sheet.add_table(table)
                sheet.auto_filter.ref = None
            else:
                sheet.auto_filter.ref = "A4:H4"
            sheet.print_area = f"A1:H{sheet.max_row}"
            atomic_save_workbook(book, path)
            return path, len(seen)
        finally:
            book.close()
