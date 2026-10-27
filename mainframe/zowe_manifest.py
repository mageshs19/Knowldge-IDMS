"""One maintained fetch inventory, with an offline workspace rebuild command.

Rows describe source dependencies, not conversion acceptance. Live fetch passes
its actual resolved dependency lists and per-member download observations. The
local rebuild never contacts Zowe and labels files only as locally available.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
from pathlib import Path
import re
from collections.abc import Iterable
from contextlib import ExitStack

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from mainframe.zowe_resolver import (
    ProgramDependencies, dclgen_members_for_record, expand_copy_dependencies,
    logical_text, mask_literals, resolve_dependencies,
)
from mainframe.zowe_rules import FETCH_INVENTORY_FILENAME
from mainframe.zowe_workspace import workspace_root, inventory_path
from mainframe.zowe_report_io import report_lock, atomic_save_workbook
from rules.program_kind_rules import IDMS_WRITE_VERBS

HEADERS = ("IDMS Program Name", "Copybook", "DCLGENs", "Sub Schema", "Mode")
SHEET_NAME = "Program Inventory"
HEADER_ROW = 4
SOURCE_EXTENSIONS = {"", ".cbl", ".cob", ".cobol", ".txt"}
_READY_MODE = re.compile(
    r"\bUSAGE[- ]MODE\s+(?:IS\s+)?(?:(?:EXCLUSIVE|PROTECTED|SHARED)\s+)?"
    r"(UPDATE|RETRIEVAL)\b", re.I,
)
_WRITE = re.compile(
    r"(?<![\w-])(" + "|".join((*IDMS_WRITE_VERBS, "CONNECT", "DISCONNECT"))
    + r")\s+([A-Z0-9][A-Z0-9-]*)(?![\w-])", re.I,
)
_READ = re.compile(r"(?<![\w-])(?:OBTAIN|FIND)\s+[A-Z0-9][A-Z0-9-]*", re.I)
_SQL_WRITE = re.compile(r"\bEXEC\s+SQL\s+(?:INSERT|UPDATE|DELETE|MERGE)\b", re.I)
_SQL_READ = re.compile(r"\bEXEC\s+SQL\s+(?:SELECT|DECLARE\s+\S+\s+CURSOR)\b", re.I)
_COPY_STMT = re.compile(r"\bCOPY\s+[^.]*\.", re.I)
_PREFIX_REPLACING = re.compile(
    r"REPLACING\s+(?:==:[A-Z0-9_-]+:==\s+BY\s+==[A-Z0-9_-]+==\s*)+\.$", re.I,
)


@dataclass
class InventoryRow:
    program: str
    copybooks: list[str] = field(default_factory=list)
    dclgens: list[str] = field(default_factory=list)
    subschema: str = ""
    mode: str = "Unknown"
    notes: dict[str, str] = field(default_factory=dict)

    def values(self) -> tuple[str, ...]:
        return (self.program, "\n".join(sorted(set(self.copybooks))),
                "\n".join(sorted(set(self.dclgens))), self.subschema, self.mode)


def classify_mode(
    sources: Iterable[tuple[str, str]], *, complete: bool = True,
) -> tuple[str, str]:
    """Inventory access mode from live source, never from program naming.

    Update evidence takes precedence over Retrieval. File WRITE does not imply
    database Update. Unknown is retained for absent evidence/incomplete COPY
    closure. This is static access classification, not call-graph certification.
    """
    update: list[str] = []
    retrieval: list[str] = []
    uncertain = not complete
    for name, source in sources:
        text = mask_literals(logical_text(source))
        if "\ufffd" in text or re.search(r">>\s*IF\b", text, re.I):
            uncertain = True
        for copy in _COPY_STMT.finditer(text):
            clause = re.search(r"\bREPLACING\b.*", copy.group(0), re.I | re.S)
            # Colon-delimited identifier-prefix substitution cannot create a
            # standalone database verb. General pseudo-text needs preprocessing.
            if clause and not _PREFIX_REPLACING.fullmatch(clause.group(0)):
                uncertain = True
        text = _COPY_STMT.sub(" ", text)
        for match in _READY_MODE.finditer(text):
            evidence = f"{name}: {match.group(0).strip()}"
            (update if match.group(1).upper() == "UPDATE" else retrieval).append(evidence)
        for match in _WRITE.finditer(text):
            update.append(f"{name}: {match.group(0).strip()}")
        if _SQL_WRITE.search(text):
            update.append(f"{name}: embedded SQL write")
        if _READ.search(text) or _SQL_READ.search(text):
            retrieval.append(f"{name}: database retrieval statement")
    if update:
        return "Update", "\n".join(dict.fromkeys(update)) + (
            "\nWARNING: incomplete source; Update is already proven." if uncertain else ""
        )
    if uncertain:
        return "Unknown", "WARNING: source/COPY closure is incomplete or requires preprocessing."
    if retrieval:
        return "Retrieval", "\n".join(dict.fromkeys(retrieval))
    return "Unknown", "WARNING: no database access-mode evidence in the available source."


def make_inventory_row(
    deps: ProgramDependencies, source: str, *, source_path: Path,
    copy_sources: Iterable[tuple[str, str]] = (),
    observations: dict[tuple[str, str], str] | None = None,
    origin: str = "Local workspace scan (not a mainframe fetch)",
) -> InventoryRow:
    """Build the five columns and per-cell evidence from the same fetch closure."""
    observations = observations or {}
    fragments = list(copy_sources)
    loaded = {name.upper() for name, _ in fragments}
    mode, evidence = classify_mode(
        [(deps.program, source), *fragments],
        complete=all(member in loaded for member in deps.copybooks),
    )
    notes = {
        HEADERS[0]: f"{origin}\nSource: {source_path}\n"
                    f"Source SHA-256 (UTF-8 text): {hashlib.sha256(source.encode()).hexdigest()}\n"
                    f"Observed at: {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        HEADERS[4]: evidence + "\nScope: this member and resolved ordinary COPY members; "
                    "external CALL targets are not analysed.",
    }
    for column, artifact, members in (
        (HEADERS[1], "COPYBOOK", deps.copybooks),
        (HEADERS[2], "DCLGEN", deps.dclgens),
        (HEADERS[3], "SUBSCHEMA", [deps.subschema] if deps.subschema else []),
    ):
        notes[column] = "\n".join(
            f"{member}: {observations.get((artifact, member), 'WARNING: availability not verified')}"
            for member in sorted(set(members))
        ) or "No dependency found in the scanned source."
    notes[HEADERS[2]] += (
        "\nDCLGENs are fetch dependencies: LR elements come from the subschema; "
        "otherwise the existing record-to-member naming rule is used. "
        "This is not the minimal SQL include list of the converted program."
    )
    if not deps.subschema:
        notes[HEADERS[3]] = (
            "No subschema reference detected in the program/resolved COPY source; "
            "left blank, no subschema lookup requested. This is not proof of runtime requirements."
        )
        if not all(member in loaded for member in deps.copybooks):
            notes[HEADERS[3]] = "WARNING: unresolved COPY source; subschema requirements are not established."
    return InventoryRow(deps.program, list(deps.copybooks), list(deps.dclgens),
                        deps.subschema, mode, notes)


def _existing_rows(path: Path) -> dict[str, InventoryRow]:
    if not path.exists():
        return {}
    try:
        book = load_workbook(path)
        try:
            if book.sheetnames != [SHEET_NAME]:
                raise ValueError("expected a single Program Inventory sheet")
            sheet = book[SHEET_NAME]
            if tuple(sheet.cell(HEADER_ROW, c).value for c in range(1, 6)) != HEADERS:
                raise ValueError("inventory column schema does not match")
            result: dict[str, InventoryRow] = {}
            for cells in sheet.iter_rows(min_row=HEADER_ROW + 1, max_col=5):
                if not any(c.value is not None for c in cells):
                    continue
                values = [str(c.value or "") for c in cells]
                name = values[0].strip().upper()
                if not name or name in result or any(c.data_type == "f" for c in cells):
                    raise ValueError("duplicate/empty member or formula in inventory")
                result[name] = InventoryRow(
                    name, values[1].splitlines(), values[2].splitlines(), values[3],
                    values[4], {HEADERS[i]: c.comment.text for i, c in enumerate(cells)
                                if c.comment},
                )
            return result
        finally:
            book.close()
    except Exception as exc:
        raise ValueError(f"Cannot read inventory {path}; original file retained: {exc}") from exc


def _workbook(rows: list[InventoryRow]) -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "IDMS | Program dependency inventory"
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17324D")
    sheet.row_dimensions[1].height = 36
    sheet.merge_cells("A2:E2")
    sheet["A2"] = (f"{len(rows)} program(s) | Maintained by fetch; partial fetches update selected rows. "
                   "Red/orange cells: unresolved evidence (see cell comments). "
                   "Mode is source access classification, not conversion approval.")
    sheet["A2"].font = Font(name="Calibri", size=10, color="626C76")
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34
    for col, title in enumerate(HEADERS, 1):
        cell = sheet.cell(HEADER_ROW, col, title)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[HEADER_ROW].height = 28
    for number, row in enumerate(rows, HEADER_ROW + 1):
        for column, value in enumerate(row.values(), 1):
            if len(value) > 32767:
                raise ValueError(f"Excel cell limit exceeded for {row.program}; no truncated report written")
            cell = sheet.cell(number, column, value)
            cell.data_type = "s"  # Member names/source text must never become Excel formulas.
            cell.font = Font(name="Calibri", size=11,
                             color="7030A0" if column == 5 else "177245")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            note = row.notes.get(HEADERS[column - 1], "")
            if note:
                if len(note) > 32767:
                    raise ValueError(f"Excel evidence limit exceeded for {row.program}")
                cell.comment = Comment(note, "IDMS fetch inventory")
            if "WARNING:" in note:
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
                cell.font = Font(name="Calibri", size=11, color="9C0006")
            elif number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F8")
        height = max(len(v.splitlines()) for v in row.values())
        sheet.row_dimensions[number].height = min(240, max(32, 16 * height + 10))
    for column, width in zip("ABCDE", (25, 31, 34, 22, 20)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B5"
    if rows:
        table = Table(displayName="IDMSPrograms", ref=f"A4:E{4 + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        sheet.add_table(table)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:E{max(4, 4 + len(rows))}"
    return book


def write_inventory(rows: Iterable[InventoryRow], path: Path | None = None) -> tuple[Path, int]:
    """Upsert programs into reports/, migrating legacy Review inventory safely.

    Both files are locked during migration. Conflicting pre-existing snapshots
    are refused rather than guessed; a failed save never removes the old file.
    """
    path = Path(path) if path is not None else inventory_path()
    legacy = path.parent.parent / "Review" / FETCH_INVENTORY_FILENAME
    migrate = (path.name == FETCH_INVENTORY_FILENAME and path.parent.name == "reports"
               and legacy.is_file() and not legacy.is_symlink() and not legacy.parent.is_symlink())
    with ExitStack() as stack:
        for candidate in sorted({path, legacy} if migrate else {path}, key=str):
            stack.enter_context(report_lock(candidate))
        records = _existing_rows(path)
        if migrate:
            for key, row in _existing_rows(legacy).items():
                if key in records and records[key] != row:
                    raise ValueError(f"Conflicting legacy/current inventory row {key}; both originals retained")
                records.setdefault(key, row)
        incoming: set[str] = set()
        for row in rows:
            key = row.program.strip().upper()
            if not key or key in incoming:
                raise ValueError("Each incoming inventory member must have one unique nonempty name")
            incoming.add(key)
            row.program = key
            records[key] = row
        book = _workbook([records[key] for key in sorted(records)])
        try:
            atomic_save_workbook(book, path)
        finally:
            book.close()
        if migrate:
            legacy.unlink()
        return path, len(records)


def migrate_legacy_inventory(root: Path | None = None) -> Path | None:
    """Local-only migration, including before conversion cleanup; no fetch call."""
    root = Path(root) if root is not None else workspace_root()
    old = root / "Review" / FETCH_INVENTORY_FILENAME
    if old.is_file() and not old.is_symlink() and not old.parent.is_symlink():
        return write_inventory((), root / "reports" / FETCH_INVENTORY_FILENAME)[0]
    return None


def _local_members(folder: Path) -> dict[str, Path]:
    result: dict[str, Path] = {}
    if not folder.is_dir():
        return result
    for path in sorted(folder.iterdir()):
        if path.is_file() and path.suffix.lower() in SOURCE_EXTENSIONS | {".cpy"}:
            key = path.stem.upper()
            if key in result:
                raise ValueError(f"Ambiguous local member {key} in {folder}")
            result[key] = path
    return result


def build_inventory_rows(root: Path | None = None) -> list[InventoryRow]:
    """Rebuild current local inputs without credentials or network access."""
    root = Path(root) if root is not None else workspace_root()
    programs = _local_members(root / "Program")
    copies = _local_members(root / "Copybook")
    schemas = _local_members(root / "Subschema")
    dclgens = _local_members(root / "DCLGen")
    rows: list[InventoryRow] = []
    for program, path in sorted(programs.items()):
        observations: dict[tuple[str, str], str] = {}

        def read_copy(member: str) -> str | None:
            target = copies.get(member)
            if target is None:
                observations[("COPYBOOK", member)] = "WARNING: not present in local Copybook folder"
                return None
            try:
                text = target.read_text(encoding="utf-8")
            except (OSError, UnicodeError):
                observations[("COPYBOOK", member)] = "WARNING: local member is unreadable"
                return None
            observations[("COPYBOOK", member)] = f"Local file: {target.relative_to(root)}"
            return text

        try:
            source = path.read_text(encoding="utf-8")
            deps = resolve_dependencies(program, source)
            fragments = expand_copy_dependencies(deps, read_copy)
            schema_path = schemas.get(deps.subschema)
            if deps.subschema:
                observations[("SUBSCHEMA", deps.subschema)] = (
                    f"Local file: {schema_path.relative_to(root)}" if schema_path
                    else "WARNING: not present; LR element resolution may be incomplete"
                )
            members = list(deps.explicit_dclgens)
            for record in deps.records:
                members.extend(dclgen_members_for_record(record, schema_path))
            deps.dclgens = sorted(set(members))
            for member in deps.dclgens:
                target = dclgens.get(member)
                observations[("DCLGEN", member)] = (
                    f"Local file: {target.relative_to(root)}" if target
                    else "WARNING: not present in local DCLGen folder"
                )
            rows.append(make_inventory_row(deps, source, source_path=path.relative_to(root),
                                           copy_sources=fragments, observations=observations))
        except (OSError, UnicodeError, ValueError) as exc:
            rows.append(InventoryRow(program, notes={
                HEADERS[0]: f"WARNING: local source/dependency scan failed ({type(exc).__name__}).",
                HEADERS[4]: "WARNING: mode not established; inspect source/dependencies.",
            }))
    return rows


def build_manifest_rows() -> list[str]:
    """Compatibility view of the inventory; no broken legacy constant imports."""
    return [" - ".join(value.replace("\n", ", ") or "[]" for value in row.values())
            for row in build_inventory_rows()]


def write_execution_manifest() -> tuple[Path | None, int]:
    """Compatibility entrypoint now maintains the same Excel inventory."""
    rows = build_inventory_rows()
    return write_inventory(rows) if rows else (None, 0)


def main() -> None:
    parser = argparse.ArgumentParser(description="Rebuild the dependency Excel inventory from local inputs; no Zowe calls.")
    parser.add_argument("--workspace", type=Path, default=None)
    args = parser.parse_args()
    root = args.workspace if args.workspace is not None else workspace_root()
    rows = build_inventory_rows(root)
    if not rows:
        print("No local program sources; existing inventory unchanged.")
        return
    path, count = write_inventory(rows, root / "reports" / FETCH_INVENTORY_FILENAME)
    print(f"Inventory: {path} ({count} programs; local availability, not live fetch verification)")


if __name__ == "__main__":
    main()
