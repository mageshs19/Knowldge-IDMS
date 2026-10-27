"""Generic fetch inventory tests: no mainframe credentials or network required."""
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from mainframe.zowe_fetcher import ZoweFetcher
from mainframe.zowe_manifest import (
    HEADERS, InventoryRow, build_inventory_rows, classify_mode, write_inventory,
    build_manifest_rows, write_execution_manifest,
)
from mainframe.zowe_resolver import resolve_dependencies
from mainframe.zowe_workspace import clean_generated, ensure_workspace, inventory_path
from mainframe.zowe_rules import ARTIFACT_EXTENSIONS, FETCH_INVENTORY_FILENAME


def fixed(body, indicator=" ", number=100):
    return f"{number:06d}{indicator}{body:<65}{number:08d}"


def read_rows(path):
    book = load_workbook(path)
    try:
        return {r[0].value: [c.value for c in r] for r in book.active.iter_rows(min_row=5, max_col=5)
                if r[0].value}
    finally:
        book.close()


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("ZOWE_WORKSPACE_ROOT", str(tmp_path))
    ensure_workspace()
    return tmp_path


@pytest.fixture
def fake_libraries(monkeypatch):
    registry = {}
    calls = []

    class Library:
        def __init__(self, artifact):
            self.artifact = artifact
            self.label = artifact
            self.libraries = [f"TEST.{artifact}"]
            self.search_history = {}

        def download(self, member):
            calls.append((self.artifact, member))
            content = registry.get((self.artifact, member))
            from mainframe.zowe_search_history import http_attempt
            self.search_history[member] = [http_attempt(self.libraries[0], member, 200 if content is not None else 404)]
            return (self.libraries[0], content.encode()) if content is not None else ("", b"")

    monkeypatch.setattr("mainframe.zowe_fetcher.LibrarySearch", Library)
    monkeypatch.setattr("mainframe.zowe_fetcher.application_code", lambda: "ZZ")
    return registry, calls


@pytest.mark.parametrize("source,expected", [
    ("READY AREA-A USAGE-MODE IS RETRIEVAL.\nWRITE FILE-ROW.", "Retrieval"),
    ("READY USAGE-MODE IS UPDATE.", "Update"),
    ("IF FLAG = 1 THEN MODIFY ABCD ELSE CONTINUE END-IF.", "Update"),
    ("STORE ABCD WHERE INSERT-ROW.", "Update"),
    ("ERASE ABCD.", "Update"),
    ("CONNECT ABCD TO SOME-SET.", "Update"),
    ("DISCONNECT ABCD FROM SOME-SET.", "Update"),
    ("OBTAIN NEXT ABCD.\nWRITE OUTPUT-RECORD.", "Retrieval"),
    ("EXEC SQL UPDATE TAB SET COL=1 END-EXEC.", "Update"),
    ("EXEC SQL SELECT COL INTO :HV FROM TAB END-EXEC.", "Retrieval"),
    ("* STORE ABCD.\nOBTAIN FIRST ABCD.", "Retrieval"),
    (fixed("MODIFY ABCD.", "*") + "\n" + fixed("OBTAIN NEXT ABCD."), "Retrieval"),
    (fixed("STORE ABCD.", "/") + "\n" + fixed("OBTAIN NEXT ABCD."), "Retrieval"),
    ("DISPLAY 'MODIFY ABCD'.\nOBTAIN NEXT ABCD. *> STORE ABCD.", "Retrieval"),
    ("DISPLAY \"STORE ABCD\".\nOBTAIN NEXT ABCD.", "Retrieval"),
    ("DISPLAY 'DON''T MODIFY ABCD'.\nOBTAIN NEXT ABCD.", "Retrieval"),
    ("MOVE 'USAGE-MODE IS UPDATE' TO MSG.\nOBTAIN NEXT ABCD.", "Retrieval"),
    ("INITIALIZE RECORD-A REPLACING ALPHANUMERIC BY SPACES.\nOBTAIN NEXT ABCD.", "Retrieval"),
    ("COPY CP1 REPLACING ==:TAG:== BY ==NEW==.\nOBTAIN NEXT ABCD.", "Retrieval"),
    ("COPY CP1 REPLACING ==OBTAIN== BY ==MODIFY==.\nOBTAIN NEXT ABCD.", "Unknown"),
    ("", "Unknown"),
    ("WRITE RECORD-A.", "Unknown"),
    ("IDENTIFICATION DIVISION.\nAUTHOR. D'ANGELO.\nPROCEDURE DIVISION.\nMODIFY ABCD.", "Update"),
])
def test_mode(source, expected):
    assert classify_mode([("ANYNAME", source)])[0] == expected


def test_missing_copy_never_certifies_retrieval():
    assert classify_mode([("P", "OBTAIN NEXT ABCD.")], complete=False)[0] == "Unknown"
    assert classify_mode([("P", "MODIFY ABCD.")], complete=False)[0] == "Update"


def test_literal_and_comment_dependency_noise():
    source = '\n'.join([
        fixed("COPY OMIT.", "*"), fixed("COPY PAGE.", "/"),
        fixed('    DISPLAY "OBTAIN FAKE COPY BAD DB BAD WITHIN BAD".'),
        fixed("    COPY 'ACTUAL'. *> COPY NOPE."),
        fixed("    COPY IDMS LR TESTLR."),
        fixed("    EXEC SQL INCLUDE EXPLICIT END-EXEC."),
        fixed("    EXEC SQL INCLUDE SQLCA END-EXEC."),
        fixed("DB TESTSUB WITHIN SCHEMA1."),
    ])
    deps = resolve_dependencies("P1", source)
    assert deps.copybooks == ["ACTUAL"]
    assert deps.records == ["TESTLR"]
    assert deps.explicit_dclgens == ["EXPLICIT"]
    assert deps.subschema == "TESTSUB"


def test_numbered_and_unsequenced_fixed_comments():
    deps = resolve_dependencies("X", "      *COPY BAD.\n       COPY GOOD.\n      /COPY BAD2.")
    assert deps.copybooks == ["GOOD"]


def test_writer_one_sheet_idempotent_upsert(workspace):
    first = InventoryRow("P1", ["C2", "C1", "C1"], ["DG1"], "SUB1", "Retrieval",
                         {HEADERS[0]: "source retained"})
    path, count = write_inventory([first, InventoryRow("P2", mode="Update")])
    assert count == 2
    write_inventory([InventoryRow("p2", ["C3"], mode="Update")])
    write_inventory([InventoryRow("p2", ["C3"], mode="Update")])
    assert read_rows(path) == {"P1": ["P1", "C1\nC2", "DG1", "SUB1", "Retrieval"],
                               "P2": ["P2", "C3", None, None, "Update"]}
    book = load_workbook(path)
    assert book.sheetnames == ["Program Inventory"]
    assert tuple(c.value for c in book.active[4]) == HEADERS
    assert book.active["A5"].comment.text == "source retained"
    assert book.active.freeze_panes == "B5"
    assert not book.active.auto_filter.ref  # Table owns the filter.
    book.close()


def test_failure_does_not_overwrite_previous_file(workspace, monkeypatch):
    path, _ = write_inventory([InventoryRow("P1")])
    before = path.read_bytes()
    monkeypatch.setattr("mainframe.zowe_report_io.os.replace", Mock(side_effect=PermissionError("locked")))
    with pytest.raises(PermissionError):
        write_inventory([InventoryRow("P2")])
    assert path.read_bytes() == before
    assert not list(path.parent.glob(".report_*.xlsx"))
    assert not path.with_suffix(".xlsx.lock").exists()


def test_corrupt_inventory_is_not_silently_replaced(workspace):
    path = inventory_path()
    path.write_bytes(b"not an Excel file")
    with pytest.raises(ValueError, match="original file retained"):
        write_inventory([InventoryRow("P1")])
    assert path.read_bytes() == b"not an Excel file"


def test_concurrent_writer_refused(workspace):
    inventory_path().with_suffix(".xlsx.lock").touch()
    with pytest.raises(RuntimeError, match="parallel"):
        write_inventory([InventoryRow("P1")])


def test_duplicate_rows_rejected(workspace):
    with pytest.raises(ValueError, match="unique"):
        write_inventory([InventoryRow("P1"), InventoryRow("p1")])


def test_excel_formula_text_is_not_executed(workspace):
    path, _ = write_inventory([InventoryRow("=EVIL", ["=1+1"])])
    book = load_workbook(path)
    assert book.active["A5"].data_type == "s"
    assert book.active["B5"].data_type == "s"
    book.close()


def test_cleanup_preserves_only_inventory_and_inputs(workspace):
    path, _ = write_inventory([InventoryRow("P1")])
    before = path.read_bytes()
    for name in ("Output/stale.cbl", "Review/approved.txt", "Program/original.cbl"):
        (workspace / name).write_text("unchanged")
    (workspace / "Review/old").mkdir()
    (workspace / "Review/old/report.txt").touch()
    clean_generated()
    assert path.read_bytes() == before
    assert list((workspace / "Output").iterdir()) == []
    assert list((workspace / "Review").iterdir()) == []
    assert path.parent == workspace / "reports"
    assert (workspace / "Program/original.cbl").read_text() == "unchanged"


def test_fetch_uses_lr_elements_and_records_missing_members(workspace, fake_libraries):
    registry, calls = fake_libraries
    registry.update({
        ("PROGRAM", "APP0001"): "COPY CP1.\nDB SUB1 WITHIN SCHEMA1.\nOBTAIN NEXT LOGREC.\nREADY USAGE-MODE IS RETRIEVAL.",
        ("COPYBOOK", "CP1"): "COPY CP2.",
        ("COPYBOOK", "CP2"): "COPY CP1.\n01 FIELD-A PIC X.",
        ("SUBSCHEMA", "SUB1"): "ADD LOGICAL RECORD LOGREC\nELEMENTS ARE XXABCD XXEFGH\n.",
        ("DCLGEN", "DZABCDTV"): "01 DCL-A PIC X.",
    })
    result = ZoweFetcher(["APP0001"]).fetch_all()
    assert result["inventory_programs"] == 1
    assert read_rows(result["inventory_path"])["APP0001"] == [
        "APP0001", "CP1\nCP2", "DZABCDTV\nDZEFGHTV", "SUB1", "Retrieval"]
    assert calls.count(("COPYBOOK", "CP1")) == 1  # Cycle terminates.
    assert ("DCLGEN", "DZGRECTV") not in calls  # Not a logical-record name guess.
    book = load_workbook(result["inventory_path"])
    assert "WARNING:" in book.active["C5"].comment.text
    assert "DZEFGHTV" in book.active["C5"].comment.text
    book.close()


def test_missing_copy_and_stale_disk_file_are_not_success(workspace, fake_libraries):
    registry, _ = fake_libraries
    registry[("PROGRAM", "P1")] = "COPY MISSING.\nOBTAIN FIRST XXABCD."
    (workspace / "Copybook/MISSING.cpy").write_text("MODIFY XXABCD.")
    result = ZoweFetcher(["P1"]).fetch_all()
    assert read_rows(result["inventory_path"])["P1"][4] == "Unknown"
    book = load_workbook(result["inventory_path"])
    assert "WARNING:" in book.active["B5"].comment.text
    book.close()


def test_failed_program_and_unattempted_selected_rows_replace_stale_success(workspace, fake_libraries):
    registry, _ = fake_libraries
    write_inventory([InventoryRow("P2", mode="Retrieval"), InventoryRow("P3", mode="Retrieval"),
                     InventoryRow("OLD", mode="Update")])
    registry[("PROGRAM", "P1")] = "OBTAIN NEXT XXABCD."
    with pytest.raises(ValueError):
        ZoweFetcher(["P1", "P2", "P3"]).fetch_all()
    rows = read_rows(inventory_path())
    assert rows["P1"][4] == "Retrieval"
    assert rows["P2"][4] == rows["P3"][4] == "Unknown"
    assert rows["OLD"][4] == "Update"


def test_write_inside_copy_is_update(workspace, fake_libraries):
    registry, _ = fake_libraries
    registry[("PROGRAM", "P1")] = "COPY CP1.\nREADY USAGE-MODE IS RETRIEVAL."
    registry[("COPYBOOK", "CP1")] = "MODIFY XXABCD."
    result = ZoweFetcher(["P1"]).fetch_all()
    assert read_rows(result["inventory_path"])["P1"][4] == "Update"


def test_partial_fetch_keeps_other_rows_and_deduplicates_selection(workspace, fake_libraries):
    registry, _ = fake_libraries
    registry[("PROGRAM", "P1")] = "OBTAIN NEXT XXABCD."
    registry[("PROGRAM", "P2")] = "MODIFY XXABCD."
    ZoweFetcher(["P1"]).fetch_all()
    ZoweFetcher(["p2", "P2"]).fetch_all()
    assert set(read_rows(inventory_path())) == {"P1", "P2"}


def test_75_programs_one_inventory_shared_downloads_once(workspace, fake_libraries):
    registry, calls = fake_libraries
    members = [f"ZZ{i:06d}" for i in range(75)]
    for member in members:
        registry[("PROGRAM", member)] = "COPY COMMON.\nDB SUB1 WITHIN SCHEMA1.\nOBTAIN NEXT XXABCD."
    registry[("COPYBOOK", "COMMON")] = "01 FIELD-A PIC X."
    registry[("SUBSCHEMA", "SUB1")] = "NO LOGICAL RECORDS"
    registry[("DCLGEN", "DZABCDTV")] = "01 DCL-A PIC X."
    result = ZoweFetcher(members).fetch_all()
    assert result["inventory_programs"] == 75
    assert len(read_rows(inventory_path())) == 75
    assert calls.count(("COPYBOOK", "COMMON")) == 1
    assert calls.count(("DCLGEN", "DZABCDTV")) == 1
    assert len(list((workspace / "reports").glob("*.xlsx"))) == 2


def test_six_source_assets_offline_match_fetch_resolution(workspace, fake_libraries):
    source_root = Path(__file__).resolve().parents[1] / "mainframe/workspace"
    registry, _ = fake_libraries
    for folder, artifact in (("Program", "PROGRAM"), ("Copybook", "COPYBOOK"),
                             ("Subschema", "SUBSCHEMA"), ("DCLGen", "DCLGEN")):
        for path in (source_root / folder).iterdir():
            if path.is_file():
                registry[(artifact, path.stem.upper())] = path.read_text(encoding="utf-8")
    members = sorted(key[1] for key in registry if key[0] == "PROGRAM")
    result = ZoweFetcher(members).fetch_all()
    live_rows = read_rows(result["inventory_path"])
    local = build_inventory_rows(workspace)
    assert len(local) == len(live_rows) == 6
    assert {r.program: list(r.values()) for r in local} == live_rows
    assert sum(r.mode == "Update" for r in local) == 2
    assert sum(r.mode == "Retrieval" for r in local) == 4
    assert len(build_manifest_rows()) == 6
    assert write_execution_manifest()[1] == 6


def test_fetch_to_convert_workbook_survives(workspace, fake_libraries, monkeypatch):
    from mainframe.zowe_pipeline_steps import step_fetch, step_convert
    registry, _ = fake_libraries
    registry[("PROGRAM", "P1")] = "OBTAIN NEXT XXABCD."
    monkeypatch.setattr("mainframe.zowe_programs.selected_programs", lambda: ["P1"])
    assert step_fetch(quiet=True).status == "OK"
    saved = inventory_path().read_bytes()
    (workspace / "Mapping Sheet/mapping.csv").write_text("A,B\n")
    monkeypatch.setattr("mainframe.zowe_pipeline_steps.run_convert", lambda: (0, 0.1))
    assert step_convert(quiet=True).status == "OK"
    assert inventory_path().read_bytes() == saved


def test_empty_program_list_is_not_default_selection(workspace):
    with pytest.raises(ValueError, match="No programs"):
        ZoweFetcher([])



def test_continued_literal_is_not_dependency_or_mode_evidence():
    source = "\n".join([
        fixed("    DISPLAY 'NOT A STATEMENT "),
        fixed("    'MODIFY FAKE COPY NOPE.'", "-"),
        fixed("    OBTAIN NEXT REALREC."),
    ])
    assert classify_mode([("ANY", source)])[0] == "Retrieval"
    deps = resolve_dependencies("ANY", source)
    assert deps.copybooks == []
    assert deps.records == ["REALREC"]
