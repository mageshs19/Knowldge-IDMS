"""Offline report/recording logic tests. Fetch, transport and network are forbidden."""
from __future__ import annotations

import ast
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock

import pytest
import requests
from openpyxl import load_workbook

from mainframe.zowe_manifest import (
    HEADERS as INVENTORY_HEADERS, InventoryRow, make_inventory_row,
    migrate_legacy_inventory, write_inventory,
)
from mainframe.zowe_resolver import resolve_dependencies
from mainframe.zowe_search_history import (
    HEADERS, HistoryRow, SearchAttempt, SearchHistoryRecorder,
    exception_attempt, http_attempt, write_search_history,
)
from mainframe.zowe_workspace import (
    clean_generated, ensure_workspace, inventory_path, reports_dir, search_history_path,
)


@pytest.fixture(autouse=True)
def offline_only(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Fetch/network prohibited: report logic tests only")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)
    from mainframe.zowe_libraries import LibrarySearch
    from mainframe.zowe_fetcher import ZoweFetcher
    for method in ("_lookup", "download", "locate", "list_members"):
        monkeypatch.setattr(LibrarySearch, method, forbidden)
    for method in ("fetch_all", "_fetch_program", "_fetch_dependencies", "_fetch_member"):
        monkeypatch.setattr(ZoweFetcher, method, forbidden)


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("ZOWE_WORKSPACE_ROOT", str(tmp_path))
    ensure_workspace()
    return tmp_path


def observation(program="P1", member="COMMON", run_id="run-1", result="Found", library="TEST.CPY(COMMON)"):
    return HistoryRow(program, "Copybook", member, library, result,
                      "Offline unit-test fixture, not a mainframe observation.", run_id,
                      "2026-10-28T00:00:00.000001+00:00")


def read_history(path):
    book = load_workbook(path)
    try:
        return [tuple(c.value or "" for c in row) for row in
                book.active.iter_rows(min_row=5, max_col=8) if row[0].value]
    finally:
        book.close()


def test_report_paths_and_workspace_override(workspace):
    assert reports_dir() == workspace / "reports"
    assert inventory_path().parent == reports_dir()
    assert search_history_path().parent == reports_dir()
    assert reports_dir().is_dir()


def test_no_subschema_keeps_empty_cell_and_no_false_missing_warning(workspace):
    source = "PROGRAM-ID. P1.\nCOPY COMMON.\nEXEC SQL SELECT X INTO :Y FROM T END-EXEC."
    row = make_inventory_row(resolve_dependencies("P1", source), source,
                             source_path=Path("Program/P1.cbl"),
                             copy_sources=[("COMMON", "01 X PIC X.")])
    assert row.subschema == ""
    assert "WARNING:" not in row.notes[INVENTORY_HEADERS[3]]
    path, _ = write_inventory([row])
    book = load_workbook(path)
    assert book.active["D5"].value is None
    book.close()


def test_unresolved_copy_does_not_claim_subschema_not_required():
    source = "COPY MISSING.\nOBTAIN NEXT REC1."
    row = make_inventory_row(resolve_dependencies("P1", source), source, source_path=Path("P1.cbl"))
    assert row.subschema == ""
    assert "WARNING:" in row.notes[INVENTORY_HEADERS[3]]
    assert row.mode == "Unknown"


def test_missing_referenced_schema_retains_its_name():
    source = "DB SUB1 WITHIN SCHEMA1.\nOBTAIN NEXT REC1."
    row = make_inventory_row(resolve_dependencies("P1", source), source, source_path=Path("P1.cbl"),
                             observations={("SUBSCHEMA", "SUB1"): "WARNING: not fetched"})
    assert row.subschema == "SUB1"
    assert "WARNING:" in row.notes[INVENTORY_HEADERS[3]]


@pytest.mark.parametrize("status,result", [(200, "Found"), (404, "Not Found"), (400, "Error"),
    (401, "Error"), (403, "Error"), (408, "Error"), (429, "Error"), (500, "Error"), (503, "Error"), (302, "Error")])
def test_status_classification(status, result):
    attempt = http_attempt("TEST.CPY", "MEM1", status)
    assert attempt.result == result
    assert attempt.library_path == "TEST.CPY(MEM1)"
    assert f"HTTP {status}" in attempt.remarks


@pytest.mark.parametrize("exc", [requests.Timeout("password=SECRET"),
    requests.ConnectionError("https://user:SECRET@example.org"),
    requests.exceptions.SSLError("token SECRET"), ValueError("SECRET"), RuntimeError("SECRET")])
def test_error_details_never_persist_raw_credentials(exc):
    attempt = exception_attempt("TEST.CPY", "MEM1", exc)
    assert attempt.result == "Error"
    assert type(exc).__name__ in attempt.remarks
    assert "SECRET" not in attempt.remarks


def test_failed_first_library_then_found_preserves_both():
    recorder = SearchHistoryRecorder(run_id="run-1")
    recorder.record("p1", "COPYBOOK", "common", [http_attempt("FIRST.CPY", "COMMON", 404),
                                                 http_attempt("SECOND.CPY", "COMMON", 200)])
    assert [(r.program, r.library_path, r.result) for r in recorder.rows] == [
        ("P1", "FIRST.CPY(COMMON)", "Not Found"), ("P1", "SECOND.CPY(COMMON)", "Found")]


def test_shared_dependency_repeated_for_each_program_with_cache_provenance():
    recorder = SearchHistoryRecorder(run_id="run-1")
    trace = [http_attempt("FIRST.CPY", "COMMON", 404), http_attempt("SECOND.CPY", "COMMON", 200)]
    recorder.record("P1", "COPYBOOK", "COMMON", trace)
    recorder.record("P2", "COPYBOOK", "COMMON", trace, cached=True)
    recorder.record("P2", "COPYBOOK", "COMMON", trace, cached=True)
    assert len(recorder.rows) == 4
    assert recorder.rows[2].program == "P2"
    assert "no additional mainframe request" in recorder.rows[2].remarks
    assert "Reused" not in recorder.rows[0].remarks


def test_shared_negative_result_is_not_hidden_for_second_program():
    recorder = SearchHistoryRecorder(run_id="run-1")
    trace = [http_attempt("ONLY.DCL", "MISSING", 404)]
    recorder.record("P1", "DCLGEN", "MISSING", trace)
    recorder.record("P2", "DCLGEN", "MISSING", trace, cached=True)
    assert [r.result for r in recorder.rows] == ["Not Found", "Not Found"]


def test_all_three_dependency_types_and_separate_members():
    recorder = SearchHistoryRecorder(run_id="run-1")
    for artifact, member in (("COPYBOOK", "CP1"), ("COPYBOOK", "CP2"),
                              ("DCLGEN", "DCL1"), ("SUBSCHEMA", "SUB1")):
        recorder.record("P1", artifact, member, [http_attempt("TEST.LIB", member, 200)])
    assert [r.artifact_type for r in recorder.rows] == ["Copybook", "Copybook", "DCLGENs", "Sub Schema"]
    assert [r.member for r in recorder.rows] == ["CP1", "CP2", "DCL1", "SUB1"]


def test_no_schema_no_invented_history_row(workspace):
    path, count = write_search_history([])
    assert count == 0
    assert read_history(path) == []
    book = load_workbook(path)
    assert tuple(c.value for c in book.active[4]) == HEADERS
    assert book.active.max_row == 4
    book.close()


def test_absent_trace_is_not_fabricated_as_found():
    recorder = SearchHistoryRecorder()
    recorder.record("P1", "COPYBOOK", "CP1", [])
    assert recorder.rows[0].result == "Not Searched"
    assert recorder.rows[0].library_path == ""


def test_local_save_error_keeps_remote_success():
    recorder = SearchHistoryRecorder()
    recorder.record("P1", "COPYBOOK", "CP1", [http_attempt("TEST.CPY", "CP1", 200)])
    recorder.issue("P1", "COPYBOOK", "CP1", "TEST.CPY(CP1)", "PermissionError: local save failed.")
    assert [r.result for r in recorder.rows] == ["Found", "Error"]


def test_history_append_new_runs_and_idempotent_republish(workspace):
    first = observation()
    path, count = write_search_history([first])
    assert count == 1
    assert write_search_history([first])[1] == 1
    second = replace(first, run_id="run-2", result="Not Found")
    assert write_search_history([second])[1] == 2
    assert [r[6] for r in read_history(path)] == ["run-1", "run-2"]
    book = load_workbook(path)
    assert book.active.tables["DependencySearches"].ref == "A4:H6"
    assert book.active.tables["DependencySearches"].autoFilter.ref == "A4:H6"
    assert book.active.auto_filter.ref is None
    book.close()


def test_history_text_not_excel_formulas(workspace):
    row = replace(observation(), remarks="=HYPERLINK(\"bad\")")
    path, _ = write_search_history([row])
    book = load_workbook(path)
    assert book.active["F5"].data_type == "s"
    book.close()


def test_history_append_failure_retains_previous_file(workspace, monkeypatch):
    path, _ = write_search_history([observation()])
    before = path.read_bytes()
    monkeypatch.setattr("mainframe.zowe_report_io.os.replace", Mock(side_effect=PermissionError()))
    with pytest.raises(PermissionError):
        write_search_history([observation(run_id="run-2")])
    assert path.read_bytes() == before
    assert not path.with_suffix(".xlsx.lock").exists()
    assert not list(path.parent.glob(".report_*.xlsx"))


def test_history_lock_refuses_parallel_writer(workspace):
    path = search_history_path()
    lock = path.with_suffix(".xlsx.lock")
    lock.touch()
    with pytest.raises(RuntimeError, match="parallel"):
        write_search_history([observation()])
    assert lock.exists()


def test_corrupt_history_is_not_replaced(workspace):
    path = search_history_path()
    path.write_bytes(b"not an xlsx")
    with pytest.raises(Exception):
        write_search_history([observation()])
    assert path.read_bytes() == b"not an xlsx"


def test_modified_history_header_is_not_overwritten(workspace):
    path, _ = write_search_history([observation()])
    book = load_workbook(path)
    book.active["A4"] = "WRONG"
    book.save(path)
    book.close()
    before = path.read_bytes()
    with pytest.raises(ValueError, match="columns"):
        write_search_history([observation(run_id="run-2")])
    assert path.read_bytes() == before


@pytest.mark.parametrize("change", [{"remarks": "x" * 32768}, {"result": "Success"},
    {"member": ""}, {"run_id": ""}, {"observed_at": "2026-10-28T00:00:00"},
    {"library_path": ""}])
def test_invalid_new_history_rows_rejected(workspace, change):
    with pytest.raises(ValueError):
        write_search_history([replace(observation(), **change)])
    assert not search_history_path().exists()


def test_history_symlink_is_never_written(workspace):
    target = workspace / "outside.xlsx"
    target.write_bytes(b"original")
    search_history_path().symlink_to(target)
    with pytest.raises(ValueError, match="symlink"):
        write_search_history([observation()])
    assert target.read_bytes() == b"original"


def test_excel_row_limit_preserves_existing(workspace, monkeypatch):
    path, _ = write_search_history([observation()])
    before = path.read_bytes()
    monkeypatch.setattr("mainframe.zowe_search_history.MAX_ROWS", 5)
    with pytest.raises(ValueError, match="row limit"):
        write_search_history([observation(run_id="run-2")])
    assert path.read_bytes() == before


def test_legacy_inventory_migrates_and_retains_unselected_rows(workspace):
    legacy = workspace / "Review/IDMS_Program_Inventory.xlsx"
    write_inventory([InventoryRow("OLD1"), InventoryRow("P1", mode="Retrieval")], legacy)
    path, count = write_inventory([InventoryRow("P1", mode="Update")])
    assert count == 2
    assert not legacy.exists()
    book = load_workbook(path)
    assert book.active["A5"].value == "OLD1"
    assert book.active["E6"].value == "Update"
    book.close()


def test_conflicting_legacy_reports_refused_without_deletion(workspace):
    write_inventory([InventoryRow("P1", mode="Update")])
    legacy = workspace / "Review/IDMS_Program_Inventory.xlsx"
    write_inventory([InventoryRow("P1", mode="Retrieval")], legacy)
    before = inventory_path().read_bytes()
    with pytest.raises(ValueError, match="Conflicting"):
        migrate_legacy_inventory()
    assert legacy.exists()
    assert inventory_path().read_bytes() == before


def test_migration_save_failure_leaves_legacy_intact(workspace, monkeypatch):
    legacy = workspace / "Review/IDMS_Program_Inventory.xlsx"
    write_inventory([InventoryRow("P1")], legacy)
    before = legacy.read_bytes()
    monkeypatch.setattr("mainframe.zowe_report_io.os.replace", Mock(side_effect=PermissionError()))
    with pytest.raises(PermissionError):
        migrate_legacy_inventory()
    assert legacy.read_bytes() == before
    assert not inventory_path().exists()


def test_cleanup_migrates_legacy_and_keeps_both_reports(workspace):
    legacy = workspace / "Review/IDMS_Program_Inventory.xlsx"
    write_inventory([InventoryRow("P1")], legacy)
    path, _ = write_search_history([observation()])
    history_before = path.read_bytes()
    (workspace / "Review/old_approval.txt").write_text("stale")
    (workspace / "Output/old.cbl").write_text("stale")
    clean_generated()
    assert inventory_path().exists()
    assert path.read_bytes() == history_before
    assert list((workspace / "Review").iterdir()) == []
    assert list((workspace / "Output").iterdir()) == []


def test_75_program_associations_append_new_run(workspace):
    recorder = SearchHistoryRecorder(run_id="fixture-run-1")
    trace = [http_attempt("FIRST.CPY", "COMMON", 404), http_attempt("SECOND.CPY", "COMMON", 200)]
    for i in range(75):
        recorder.record(f"P{i:07d}", "COPYBOOK", "COMMON", trace, cached=i > 0)
    path, count = write_search_history(recorder.rows)
    assert count == 150
    assert len({row[0] for row in read_history(path)}) == 75
    assert write_search_history([replace(r, run_id="fixture-run-2") for r in recorder.rows])[1] == 300


def test_publish_reports_keeps_history_if_inventory_fails(workspace, monkeypatch):
    from mainframe.zowe_fetcher import ZoweFetcher
    obj = object.__new__(ZoweFetcher)
    obj.search_history = SearchHistoryRecorder(run_id="run-1", rows=[observation()])
    monkeypatch.setattr("mainframe.zowe_manifest.write_inventory", Mock(side_effect=PermissionError()))
    with pytest.raises(RuntimeError, match="publication failed"):
        obj._publish_reports([InventoryRow("P1")])
    assert read_history(search_history_path())[0][0] == "P1"


def test_static_integration_records_attempts_before_success_return():
    """Inspect adapter integration without executing ANY lookup/fetch function."""
    root = Path(__file__).resolve().parents[1]
    tree = ast.parse((root / "mainframe/zowe_libraries.py").read_text())
    method = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "_lookup")
    loop = next(n for n in method.body if isinstance(n, ast.For))
    append_at = next(i for i, n in enumerate(loop.body) if isinstance(n, ast.Expr)
                     and isinstance(n.value, ast.Call) and ast.unparse(n.value.func) == "trace.append")
    success_at = next(i for i, n in enumerate(loop.body) if isinstance(n, ast.If)
                      and ast.unparse(n.test) == "found")
    assert append_at < success_at
    fetch_source = (root / "mainframe/zowe_fetcher.py").read_text()
    assert "self.search_history.record(self._active_program, artifact, member" in fetch_source
    assert "cached=True" in fetch_source
    assert "self._publish_reports(inventory.values())" in fetch_source
